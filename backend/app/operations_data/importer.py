from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook

MAX_IMPORT_BYTES = 10 * 1024 * 1024
HEADER_ALIASES = {
    "order_id": {"訂單編號", "訂單號", "order_id", "order id"},
    "order_date": {"訂單日期", "建立日期", "order_date", "order date"},
    "warehouse": {"倉別", "倉庫", "warehouse"},
    "customer_name": {"客戶", "客戶名稱", "customer", "customer_name"},
    "promised_at": {"承諾出貨時間", "應出貨時間", "預計出貨時間", "promised_at"},
    "completed_at": {"實際出貨時間", "出貨時間", "完成時間", "completed_at"},
    "status": {"狀態", "status"},
    "urgent": {"急單", "是否急單", "urgent"},
    "exception_count": {"異常數", "異常次數", "exception_count"},
    "processing_minutes": {"處理分鐘", "處理時間分鐘", "processing_minutes"},
    "worker_hours": {"人時", "作業人時", "worker_hours"},
}


@dataclass(frozen=True)
class ParsedOperationalOrder:
    order_id: str
    order_date: date
    warehouse: str
    customer_name: str | None
    promised_at: datetime | None
    completed_at: datetime | None
    status: str
    urgent: bool
    exception_count: int
    processing_minutes: int | None
    worker_hours: float | None


@dataclass(frozen=True)
class ParsedOperationalImport:
    records: tuple[ParsedOperationalOrder, ...]
    warnings: tuple[str, ...]


def parse_operational_file(content: bytes, suffix: str) -> ParsedOperationalImport:
    if not content:
        raise ValueError("上傳的營運報表是空白檔案。")
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError("營運報表超過 10 MB，請縮小後再匯入。")
    rows = _xlsx_rows(content) if suffix == ".xlsx" else _csv_rows(content)
    header = next(rows, None)
    if not header:
        raise ValueError("營運報表沒有欄位名稱。")
    columns = _resolve_columns(header)
    for required in ("order_id", "order_date", "warehouse"):
        if required not in columns:
            label = {"order_id": "訂單編號", "order_date": "訂單日期", "warehouse": "倉別"}[
                required
            ]
            raise ValueError(f"營運報表缺少「{label}」欄位。")

    records: dict[str, ParsedOperationalOrder] = {}
    warnings: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        order_id = _text(_value(row, columns["order_id"]))
        if not order_id:
            continue
        try:
            order_date = _date(_value(row, columns["order_date"]))
        except ValueError:
            warnings.append(f"第 {row_number} 列訂單日期無法辨識，已略過。")
            continue
        warehouse = _text(_value(row, columns["warehouse"]))
        if not warehouse:
            warnings.append(f"第 {row_number} 列沒有倉別，已略過。")
            continue
        record = ParsedOperationalOrder(
            order_id=order_id,
            order_date=order_date,
            warehouse=warehouse,
            customer_name=_optional_text(_optional(row, columns, "customer_name")),
            promised_at=_datetime_or_none(_optional(row, columns, "promised_at")),
            completed_at=_datetime_or_none(_optional(row, columns, "completed_at")),
            status=_status(_optional(row, columns, "status")),
            urgent=_boolean(_optional(row, columns, "urgent")),
            exception_count=max(0, _integer(_optional(row, columns, "exception_count")) or 0),
            processing_minutes=_integer(_optional(row, columns, "processing_minutes")),
            worker_hours=_float(_optional(row, columns, "worker_hours")),
        )
        if order_id in records:
            warnings.append(f"訂單 {order_id} 重複出現，採用最後一列。")
        records[order_id] = record
    if not records:
        raise ValueError("營運報表中沒有可匯入的訂單資料。")
    return ParsedOperationalImport(tuple(records.values()), tuple(warnings))


def _xlsx_rows(content: bytes):  # type: ignore[no-untyped-def]
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        sheet = workbook.active
        return iter(sheet.iter_rows(min_row=1, max_row=100000, max_col=40, values_only=True))
    except Exception as exc:  # noqa: BLE001
        raise ValueError("無法讀取這份 Excel，請確認檔案沒有損壞。") from exc


def _csv_rows(content: bytes):  # type: ignore[no-untyped-def]
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("big5")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV 請使用 UTF-8 或 Big5 編碼。") from exc
    return iter(csv.reader(StringIO(text)))


def _resolve_columns(header: tuple[Any, ...] | list[str]) -> dict[str, int]:
    normalized = {_text(value).casefold(): index for index, value in enumerate(header)}
    return {
        field: normalized[alias]
        for field, aliases in HEADER_ALIASES.items()
        for alias in aliases
        if alias in normalized
    }


def _value(row, index: int) -> Any:  # type: ignore[no-untyped-def]
    return row[index] if index < len(row) else None


def _optional(row, columns: dict[str, int], field: str) -> Any:  # type: ignore[no-untyped-def]
    return _value(row, columns[field]) if field in columns else None


def _text(value: Any) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def _optional_text(value: Any) -> str | None:
    result = _text(value)
    return result or None


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = _parse_datetime(_text(value))
    if parsed is None:
        raise ValueError("invalid date")
    return parsed.date()


def _datetime_or_none(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or UTC)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time(), UTC)
    return _parse_datetime(_text(value))


def _parse_datetime(value: str) -> datetime | None:
    if not value:
        return None
    normalized = value.replace("/", "-")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(normalized, fmt).replace(tzinfo=UTC)
        except ValueError:
            continue
    try:
        parsed = datetime.fromisoformat(normalized)
        return parsed.replace(tzinfo=parsed.tzinfo or UTC)
    except ValueError:
        return None


def _status(value: Any) -> str:
    text = _text(value).casefold()
    if text in {"完成", "已完成", "已出貨", "done", "completed", "shipped"}:
        return "COMPLETED"
    if text in {"取消", "已取消", "cancelled", "canceled"}:
        return "CANCELLED"
    if text in {"處理中", "作業中", "in_progress", "in progress"}:
        return "IN_PROGRESS"
    return "OPEN"


def _boolean(value: Any) -> bool:
    return _text(value).casefold() in {"1", "true", "yes", "y", "是", "急單"}


def _integer(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).replace(",", "")))
    except ValueError:
        return None


def _float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return None
