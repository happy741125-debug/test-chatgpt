from __future__ import annotations

import csv
from dataclasses import dataclass, replace
from datetime import UTC, date, datetime
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook

MAX_IMPORT_BYTES = 15 * 1024 * 1024

ORDER_ALIASES = {
    "order_id": {"訂單編號", "order_id", "order id"},
    "channel": {"通路名稱", "銷售通路", "channel"},
    "platform": {"電商平台", "platform"},
    "shipping_type": {"物流類型", "物流方式", "shipping_type"},
    "amount": {"訂單金額", "amount", "total"},
    "urgent": {"急單", "urgent"},
    "reserved_ship_date": {"預約出貨日", "reserved_ship_date"},
    "shipped_at": {"出貨時間", "shipped_at"},
    "order_status": {"訂單狀態", "狀態", "status"},
    "created_at": {"建立時間", "created_at"},
    "sku": {"品號", "sku"},
}

INVENTORY_ALIASES = {
    "merchant": {"貨主名稱", "貨主", "merchant"},
    "sku": {"品號", "sku"},
    "product_name": {"商品名稱", "品名", "product_name"},
    "inventory_type": {"庫存類型", "庫別", "inventory_type"},
    "quantity": {"數量", "quantity"},
    "batch": {"批號", "batch"},
    "expiration_date": {"效期", "有效期限", "expiration_date"},
    "status": {"狀態", "status"},
    "available": {"可用數量", "available"},
    "allocated": {"已分配數量", "allocated"},
}

OPERATIONAL_ALIASES = {
    "inbound": {
        "external_id": {"單號", "inbound_id"},
        "line_key": {"品號", "sku"},
        "occurred_on": {"實際入倉日"},
        "planned_on": {"預計入倉日期"},
        "created_on": {"建立時間"},
        "category": {"進倉類別", "category"},
        "status": {"狀態", "status"},
        "planned_quantity": {"預計入庫數量", "planned_quantity"},
        "accepted_quantity": {"實際驗收數量", "accepted_quantity"},
        "completed_quantity": {"實際上架數量", "completed_quantity"},
        "merchant": {"貨主", "貨主名稱", "merchant"},
    },
    "returns": {
        "external_id": {"退貨單號", "return_id"},
        "line_key": {"品號", "sku"},
        "occurred_on": {"退貨日期"},
        "created_on": {"建立時間"},
        "category": {"庫存類型", "category"},
        "status": {"狀態", "status"},
        "item_count": {"數量", "件數", "quantity"},
        "merchant": {"貨主", "貨主名稱", "merchant"},
        "order_id": {"訂單編號", "order_id"},
    },
    "picking": {
        "external_id": {"揀貨單編號", "picking_id"},
        "occurred_on": {"揀貨日期"},
        "created_on": {"建立時間"},
        "warehouse": {"倉庫", "warehouse"},
        "channel": {"銷售通路", "channel"},
        "status": {"狀態", "status"},
        "shipment_count": {"出貨單數", "shipment_count"},
        "item_count": {"總件數", "件數", "item_count"},
        "merchant": {"貨主", "貨主名稱", "merchant"},
    },
    "consignment": {
        "external_id": {"託運單號", "托運單號", "consignment_id"},
        "occurred_on": {"配達日"},
        "created_on": {"建立時間"},
        "category": {"類別", "物流類型", "category"},
        "channel": {"模式", "mode"},
        "status": {"狀態", "status"},
        "shipment_count": {"件數", "shipment_count"},
        "merchant": {"貨主", "貨主名稱", "merchant"},
        "order_id": {"訂單編號", "order_id"},
    },
}


@dataclass(frozen=True)
class ParsedOrder:
    order_id: str
    channel: str | None
    platform: str | None
    shipping_type: str | None
    amount: float | None
    urgent: bool
    reserved_ship_date: date | None
    shipped_at: datetime | None
    order_status: str | None
    source_created_at: datetime | None
    skus: tuple[str, ...] = ()


@dataclass(frozen=True)
class ParsedInventory:
    merchant: str
    sku: str
    product_name: str | None
    inventory_type: str | None
    quantity: int
    batch: str | None
    expiration_date: date | None
    status: str | None
    available: int | None
    allocated: int | None


@dataclass(frozen=True)
class ParsedOperationalRecord:
    source_key: str
    kind: str
    occurred_on: date | None
    category: str | None
    warehouse: str | None
    channel: str | None
    status: str | None
    planned_quantity: int
    accepted_quantity: int
    completed_quantity: int
    shipment_count: int
    item_count: int
    merchant: str | None = None
    order_id: str | None = None
    sku: str | None = None


@dataclass(frozen=True)
class ParsedImport:
    orders: tuple[ParsedOrder, ...] = ()
    inventory: tuple[ParsedInventory, ...] = ()
    operational: tuple[ParsedOperationalRecord, ...] = ()
    warnings: tuple[str, ...] = ()


def detect_import_kind(content: bytes, suffix: str) -> str:
    """Identify a supported GoWarehouse export from its header row."""
    if not content:
        raise ValueError("上傳的檔案是空白檔案。")
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError("檔案超過 15 MB，請縮小後再匯入。")
    rows = _xlsx_rows(content) if suffix == ".xlsx" else _csv_rows(content)
    header = next(rows, None)
    if not header:
        raise ValueError("檔案沒有欄位名稱。")
    normalized = {_text(value).casefold() for value in header if _text(value)}

    def has_any(names: set[str]) -> bool:
        return any(name.casefold() in normalized for name in names)

    for kind in ("returns", "picking", "consignment"):
        if has_any(OPERATIONAL_ALIASES[kind]["external_id"]):
            return kind
    if has_any(ORDER_ALIASES["order_id"]):
        return "orders"
    if has_any(OPERATIONAL_ALIASES["inbound"]["external_id"]):
        return "inbound"
    if has_any(INVENTORY_ALIASES["merchant"]) and has_any(INVENTORY_ALIASES["sku"]):
        return "inventory"
    raise ValueError(
        "無法自動判斷檔案類型；請確認這是訂單、庫存、進倉、退貨、揀貨或托運匯出檔。"
    )


def parse_orders_file(content: bytes, suffix: str) -> ParsedImport:
    rows, columns = _open(content, suffix, ORDER_ALIASES, required=("order_id",), label="訂單")
    seen: dict[str, ParsedOrder] = {}
    warnings: list[str] = []
    for _row_number, row in rows:
        order_id = _text(_get(row, columns, "order_id"))
        if not order_id:
            continue
        if order_id in seen:
            # Orders export is line-item level; keep the first (order-level fields repeat).
            sku = _opt_text(_get(row, columns, "sku"))
            if sku and sku not in seen[order_id].skus:
                seen[order_id] = replace(
                    seen[order_id], skus=(*seen[order_id].skus, sku)
                )
            continue
        sku = _opt_text(_get(row, columns, "sku"))
        seen[order_id] = ParsedOrder(
            order_id=order_id,
            channel=_opt_text(_get(row, columns, "channel")),
            platform=_opt_text(_get(row, columns, "platform")),
            shipping_type=_opt_text(_get(row, columns, "shipping_type")),
            amount=_float(_get(row, columns, "amount")),
            urgent=_boolean(_get(row, columns, "urgent")),
            reserved_ship_date=_date_or_none(_get(row, columns, "reserved_ship_date")),
            shipped_at=_datetime_or_none(_get(row, columns, "shipped_at")),
            order_status=_opt_text(_get(row, columns, "order_status")),
            source_created_at=_datetime_or_none(_get(row, columns, "created_at")),
            skus=(sku,) if sku else (),
        )
    if not seen:
        raise ValueError("訂單匯出檔中沒有可匯入的訂單資料。")
    return ParsedImport(orders=tuple(seen.values()), warnings=tuple(warnings))


def parse_inventory_file(content: bytes, suffix: str) -> ParsedImport:
    rows, columns = _open(
        content, suffix, INVENTORY_ALIASES, required=("merchant", "sku"), label="庫存"
    )
    records: list[ParsedInventory] = []
    warnings: list[str] = []
    for _row_number, row in rows:
        merchant = _text(_get(row, columns, "merchant"))
        sku = _text(_get(row, columns, "sku"))
        if not merchant or not sku:
            continue
        records.append(
            ParsedInventory(
                merchant=merchant,
                sku=sku,
                product_name=_opt_text(_get(row, columns, "product_name")),
                inventory_type=_opt_text(_get(row, columns, "inventory_type")),
                quantity=_integer(_get(row, columns, "quantity")) or 0,
                batch=_opt_text(_get(row, columns, "batch")),
                expiration_date=_date_or_none(_get(row, columns, "expiration_date")),
                status=_opt_text(_get(row, columns, "status")),
                available=_integer(_get(row, columns, "available")),
                allocated=_integer(_get(row, columns, "allocated")),
            )
        )
    if not records:
        raise ValueError("庫存匯出檔中沒有可匯入的庫存資料。")
    return ParsedImport(inventory=tuple(records), warnings=tuple(warnings))


def parse_operational_file(content: bytes, suffix: str, kind: str) -> ParsedImport:
    aliases = OPERATIONAL_ALIASES.get(kind)
    if aliases is None:
        raise ValueError("不支援的 GoWarehouse 匯出類型。")
    rows, columns = _open(
        content,
        suffix,
        aliases,
        required=("external_id",),
        label={
            "inbound": "進倉單",
            "returns": "退貨單",
            "picking": "揀貨單",
            "consignment": "托運單",
        }[kind],
    )
    records: list[ParsedOperationalRecord] = []
    source_occurrences: dict[str, int] = {}
    for row_number, row in rows:
        external_id = _text(_get(row, columns, "external_id"))
        if not external_id:
            continue
        line_key = _text(_get(row, columns, "line_key")) or str(row_number)
        source_base = f"{external_id}|{line_key}"
        occurrence = source_occurrences.get(source_base, 0) + 1
        source_occurrences[source_base] = occurrence
        records.append(
            ParsedOperationalRecord(
                source_key=f"{source_base}|{occurrence}",
                kind=kind,
                occurred_on=(
                    _date_or_none(_get(row, columns, "occurred_on"))
                    or _date_or_none(_get(row, columns, "planned_on"))
                    or _date_or_none(_get(row, columns, "created_on"))
                ),
                category=_opt_text(_get(row, columns, "category")),
                warehouse=_opt_text(_get(row, columns, "warehouse")),
                channel=_opt_text(_get(row, columns, "channel")),
                status=_opt_text(_get(row, columns, "status")),
                planned_quantity=_integer(_get(row, columns, "planned_quantity")) or 0,
                accepted_quantity=_integer(_get(row, columns, "accepted_quantity")) or 0,
                completed_quantity=_integer(_get(row, columns, "completed_quantity")) or 0,
                shipment_count=_integer(_get(row, columns, "shipment_count")) or 0,
                item_count=_integer(_get(row, columns, "item_count")) or 0,
                merchant=_opt_text(_get(row, columns, "merchant")),
                order_id=_opt_text(_get(row, columns, "order_id")),
                sku=_opt_text(_get(row, columns, "line_key")),
            )
        )
    if not records:
        raise ValueError("匯出檔中沒有可匯入的營運資料。")
    return ParsedImport(operational=tuple(records))


# --------------------------------------------------------------------------- #
# Shared parsing helpers
# --------------------------------------------------------------------------- #
def _open(
    content: bytes,
    suffix: str,
    aliases: dict[str, set[str]],
    *,
    required: tuple[str, ...],
    label: str,
):  # type: ignore[no-untyped-def]
    if not content:
        raise ValueError(f"上傳的{label}檔是空白檔案。")
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError(f"{label}檔超過 15 MB，請縮小後再匯入。")
    rows = _xlsx_rows(content) if suffix == ".xlsx" else _csv_rows(content)
    header = next(rows, None)
    if not header:
        raise ValueError(f"{label}檔沒有欄位名稱。")
    columns = _resolve_columns(header, aliases)
    for field in required:
        if field not in columns:
            names = "、".join(sorted(aliases[field]))
            raise ValueError(f"{label}檔缺少必要欄位（{names}）。")
    return enumerate(rows, start=2), columns


def _xlsx_rows(content: bytes):  # type: ignore[no-untyped-def]
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
        sheet = workbook.active
        yield from sheet.iter_rows(values_only=True)
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ValueError("無法讀取這份 Excel，請確認檔案沒有損壞。") from exc


def _csv_rows(content: bytes):  # type: ignore[no-untyped-def]
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("big5")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV 請使用 UTF-8 或 Big5 編碼。") from exc
    yield from csv.reader(StringIO(text))


def _resolve_columns(
    header: tuple[Any, ...] | list[Any], aliases: dict[str, set[str]]
) -> dict[str, int]:
    normalized = {_text(value).casefold(): index for index, value in enumerate(header)}
    resolved: dict[str, int] = {}
    for field, names in aliases.items():
        for alias in names:
            if alias.casefold() in normalized:
                resolved[field] = normalized[alias.casefold()]
                break
    return resolved


def _get(row, columns: dict[str, int], field: str) -> Any:  # type: ignore[no-untyped-def]
    index = columns.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def _text(value: Any) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def _opt_text(value: Any) -> str | None:
    return _text(value) or None


def _boolean(value: Any) -> bool:
    return _text(value).casefold() in {"1", "true", "yes", "y", "是", "急單"}


def _integer(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).replace(",", "")))
    except ValueError:
        return None


def _float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return None


def _date_or_none(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = _parse_datetime(_text(value))
    return parsed.date() if parsed else None


def _datetime_or_none(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or UTC)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time(), UTC)
    return _parse_datetime(_text(value))


def _parse_datetime(value: str) -> datetime | None:
    if not value:
        return None
    normalized = value.replace("/", "-")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(normalized, fmt).replace(tzinfo=UTC)
        except ValueError:
            continue
    try:
        parsed = datetime.fromisoformat(normalized)
        return parsed.replace(tzinfo=parsed.tzinfo or UTC)
    except ValueError:
        return None
