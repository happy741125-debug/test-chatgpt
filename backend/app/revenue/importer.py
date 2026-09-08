from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MONTH_SHEET = re.compile(r"^20\d{4}$")
CUSTOMER_LABEL = re.compile(r"^\s*(\d{4})(?:\s*[|｜_\-]?\s*)(.*)$")
KNOWN_WAREHOUSES = {"汐止", "淡水"}
MAX_IMPORT_BYTES = 10 * 1024 * 1024


@dataclass(frozen=True)
class ParsedRevenueRecord:
    period: str
    warehouse: str
    customer_code: str | None
    customer_name: str
    source_customer_label: str
    warehouse_rent: Decimal
    handling_system: Decimal
    processing: Decimal
    logistics: Decimal
    other: Decimal
    total: Decimal


@dataclass(frozen=True)
class ParsedRevenueIssue:
    period: str
    severity: str
    code: str
    message: str
    cell_reference: str | None = None
    source_value: Decimal | None = None
    calculated_value: Decimal | None = None
    difference: Decimal | None = None


@dataclass(frozen=True)
class ParsedRevenueWorkbook:
    periods: tuple[str, ...]
    records: tuple[ParsedRevenueRecord, ...]
    issues: tuple[ParsedRevenueIssue, ...]


def parse_revenue_workbook(content: bytes) -> ParsedRevenueWorkbook:
    if not content:
        raise ValueError("上傳的 Excel 是空白檔案。")
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError("Excel 檔案超過 10 MB，請縮小後再匯入。")

    try:
        # The imported workbook only needs cached values and labels. Loading both
        # formula and value copies can exhaust a small Render instance, so keep a
        # single read-only workbook in memory.
        value_book = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - return one safe import message
        raise ValueError("無法讀取這份 Excel，請確認檔案沒有損壞。") from exc

    periods = sorted(name for name in value_book.sheetnames if MONTH_SHEET.fullmatch(name))
    if not periods:
        raise ValueError("找不到以 YYYYMM 命名的月份工作表。")

    records: list[ParsedRevenueRecord] = []
    issues: list[ParsedRevenueIssue] = []
    for sheet_name in periods:
        period = f"{sheet_name[:4]}-{sheet_name[4:]}"
        sheet = value_book[sheet_name]
        rows = sheet.iter_rows(min_row=1, max_row=5000, max_col=30, values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError(f"{sheet.title} 是空白工作表。")
        total_column = _total_column(sheet.title, header)
        category_columns = _category_columns(header, total_column)
        seen: set[tuple[str, str]] = set()
        month_records: list[ParsedRevenueRecord] = []
        has_row_total_mismatch = False
        total_cell: int | None = None
        source_total: Decimal | None = None

        for row_number, row in enumerate(rows, start=2):
            warehouse = _clean_text(_cell_value(row, 1))
            customer_label = _clean_text(_cell_value(row, 2))
            if not warehouse and customer_label.endswith("總金額"):
                if customer_label == "總金額" or total_cell is None:
                    total_cell = row_number
                    source_total = _amount_or_none(_cell_value(row, total_column))
                continue
            if warehouse not in KNOWN_WAREHOUSES or not customer_label:
                continue

            amounts = {key: Decimal("0") for key in _category_keys()}
            has_amount = False
            for column, category in category_columns.items():
                amount = _amount(_cell_value(row, column))
                amounts[category] += amount
                has_amount = has_amount or amount != 0
            if not has_amount:
                continue

            customer_code, customer_name = _split_customer(customer_label)
            duplicate_key = (warehouse, customer_label)
            if duplicate_key in seen:
                issues.append(
                    ParsedRevenueIssue(
                        period=period,
                        severity="WARNING",
                        code="DUPLICATE_CUSTOMER_ROW",
                        cell_reference=f"{sheet_name}!B{row_number}",
                        message=f"{customer_label} 在同月份、同倉別出現重複明細，系統已合併計算。",
                    )
                )
                existing_index = next(
                    index
                    for index, record in enumerate(month_records)
                    if record.warehouse == warehouse
                    and record.source_customer_label == customer_label
                )
                month_records[existing_index] = _merge_records(
                    month_records[existing_index], amounts
                )
                continue

            seen.add(duplicate_key)
            total = sum(amounts.values(), Decimal("0"))
            month_records.append(
                ParsedRevenueRecord(
                    period=period,
                    warehouse=warehouse,
                    customer_code=customer_code,
                    customer_name=customer_name,
                    source_customer_label=customer_label,
                    total=total,
                    **amounts,
                )
            )

            cached_row_total = _amount_or_none(_cell_value(row, total_column))
            if cached_row_total is not None and abs(cached_row_total - total) >= Decimal("0.01"):
                has_row_total_mismatch = True
                issues.append(
                    ParsedRevenueIssue(
                        period=period,
                        severity="ERROR",
                        code="ROW_TOTAL_MISMATCH",
                        cell_reference=(
                            f"{sheet_name}!{get_column_letter(total_column)}{row_number}"
                        ),
                        message=f"{customer_label} 的 Excel 加總與明細不一致，系統已採用明細重算。",
                        source_value=cached_row_total,
                        calculated_value=total,
                        difference=total - cached_row_total,
                    )
                )

        if not month_records:
            issues.append(
                ParsedRevenueIssue(
                    period=period,
                    severity="ERROR",
                    code="NO_DETAIL_ROWS",
                    message=f"{sheet_name} 找不到可匯入的汐止或淡水客戶明細。",
                )
            )
            continue

        records.extend(month_records)
        calculated_month_total = sum((record.total for record in month_records), Decimal("0"))
        if total_cell is None:
            issues.append(
                ParsedRevenueIssue(
                    period=period,
                    severity="WARNING",
                    code="MISSING_MONTH_TOTAL",
                    message=f"{sheet_name} 找不到月份總金額，系統已用客戶明細重算。",
                    calculated_value=calculated_month_total,
                )
            )
        else:
            total_reference = (
                f"{sheet_name}!{get_column_letter(total_column)}{total_cell}"
            )
            if source_total is None:
                issues.append(
                    ParsedRevenueIssue(
                        period=period,
                        severity="WARNING",
                        code="MISSING_CACHED_TOTAL",
                        cell_reference=total_reference,
                        message=f"{sheet_name} 的總金額尚未計算，系統已用客戶明細重算。",
                        calculated_value=calculated_month_total,
                    )
                )
            elif (
                abs(source_total - calculated_month_total) >= Decimal("0.01")
                and not has_row_total_mismatch
            ):
                issues.append(
                    ParsedRevenueIssue(
                        period=period,
                        severity="ERROR",
                        code="MONTH_TOTAL_MISMATCH",
                        cell_reference=total_reference,
                        message=f"{sheet_name} 的 Excel 月總額少算或多算，系統已採用明細重算。",
                        source_value=source_total,
                        calculated_value=calculated_month_total,
                        difference=calculated_month_total - source_total,
                    )
                )

    if not records:
        raise ValueError("Excel 中沒有可匯入的營收明細。")
    imported_periods = tuple(sorted({record.period for record in records}))
    value_book.close()
    return ParsedRevenueWorkbook(imported_periods, tuple(records), tuple(issues))


def _category_keys() -> tuple[str, ...]:
    return ("warehouse_rent", "handling_system", "processing", "logistics", "other")


def _cell_value(row: tuple[object, ...], column: int) -> object | None:
    return row[column - 1] if column <= len(row) else None


def _total_column(sheet_title: str, header: tuple[object, ...]) -> int:
    for column in range(3, min(len(header), 30) + 1):
        if "總金額" in _clean_text(_cell_value(header, column)):
            return column
    raise ValueError(f"{sheet_title} 找不到總金額欄位。")


def _category_columns(header: tuple[object, ...], total_column: int) -> dict[int, str]:
    result: dict[int, str] = {}
    for column in range(3, total_column):
        label = _clean_text(_cell_value(header, column))
        if "倉租" in label:
            category = "warehouse_rent"
        elif "理貨" in label or "系統" in label:
            category = "handling_system"
        elif "加工" in label:
            category = "processing"
        elif "物流" in label or "運送" in label or "宅配" in label:
            category = "logistics"
        else:
            category = "other"
        result[column] = category
    return result


def _split_customer(label: str) -> tuple[str | None, str]:
    match = CUSTOMER_LABEL.match(label)
    if not match:
        return None, label
    code, name = match.groups()
    cleaned_name = name.strip(" _|｜-")
    return code, cleaned_name or label


def _clean_text(value: object) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def _amount(value: object) -> Decimal:
    parsed = _amount_or_none(value)
    return parsed if parsed is not None else Decimal("0")


def _amount_or_none(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    try:
        return Decimal(str(value).replace(",", "")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def _merge_records(
    existing: ParsedRevenueRecord, amounts: dict[str, Decimal]
) -> ParsedRevenueRecord:
    merged = {
        key: getattr(existing, key) + amounts[key]
        for key in _category_keys()
    }
    return ParsedRevenueRecord(
        period=existing.period,
        warehouse=existing.warehouse,
        customer_code=existing.customer_code,
        customer_name=existing.customer_name,
        source_customer_label=existing.source_customer_label,
        total=sum(merged.values(), Decimal("0")),
        **merged,
    )
